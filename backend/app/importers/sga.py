"""Leitores dedicados das fontes reais do atendimento (repositório Estatisticas-ATD):

- PDFs do Novo SGA: "Atendimentos concluídos", "Serviços executados" e
  "Tempos médios por Atendente", um arquivo por dia ("Período de d a d").
- Planilha "Guichês - Tabela Consolidada" (cadastro de quem ocupa cada guichê/sala).
- Metadados no nome do arquivo: "AAAA.MM.DD - Tipo - Relatório.ext" (data e tipo de dia).

As regras reproduzem Scripts/gerar_relatorio_atendimento.py daquele repositório:
linhas de TMA por regex no texto, tickets em blocos de 5 linhas na tabela, sala 7
(CPD) descartada e rótulos "guiche N"/"sala N" resolvidos pelo cadastro.
"""
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime

from .parse import duracao, norm

# ---------- nome do arquivo ----------
_NOME_RE = re.compile(r"^(\d{4})\.(\d{2})\.(\d{2})(?:\s*-\s*(.+?)\s*-)?")
TIPOS_NOME = {"normal": "Normal", "sabado": "Sábado", "domingo": "Domingo", "feriado": "Feriado"}


def meta_nome(nome: str) -> tuple[date | None, str | None]:
    """Data e tipo de dia do padrão de nomes dos relatórios diários (tipo é opcional)."""
    m = _NOME_RE.match(nome)
    if not m:
        return None, None
    try:
        d = date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None, None
    tipo = m[4].strip() if m[4] else None
    return d, (TIPOS_NOME.get(norm(tipo), tipo) if tipo else None)


# ---------- rótulos de guichê/sala ----------
_ROTULO_RE = re.compile(r"^(guiche|sala)\s*0*(\d+)$", re.IGNORECASE)
_RESOLVIDO_RE = re.compile(r"—\s*(guiche|sala)\s+(\d+)\s*$", re.IGNORECASE)
SALAS_IGNORADAS = {7}  # sala 7 = CPD, não é posto de atendimento


def rotulo(texto: str) -> str:
    """'guiche23', 'guiche 023', 'sala01' → 'guiche 23', 'sala 1'; nomes ficam como estão."""
    t = re.sub(r"\s+", " ", texto or "").strip()
    m = _ROTULO_RE.match(t)
    return f"{m[1].lower()} {int(m[2])}" if m else t


def origem_de(atendente: str) -> str | None:
    """Recupera o rótulo cru de um nome já resolvido ('Marta (MAPMA) — guiche 11' → 'guiche 11')."""
    m = _ROTULO_RE.match(atendente.strip()) or _RESOLVIDO_RE.search(atendente)
    return f"{m[1].lower()} {int(m[2])}" if m else None


def ignorado(rot: str) -> bool:
    m = _ROTULO_RE.match(rot)
    return bool(m and m[1].lower() == "sala" and int(m[2]) in SALAS_IGNORADAS)


@dataclass
class Ocupacao:
    tipo: str
    numero: int
    atendente: str
    grupo: str
    vig_inicio: date | None
    vig_fim: date | None


def resolvedor(cadastro: list[Ocupacao]) -> Callable[[str, date], str]:
    """Rótulo + data → 'Nome (GRUPO) — guiche N'. Um guichê cujas linhas têm todas a mesma
    pessoa vale para todo o histórico; só guichês com troca real de ocupante usam a vigência."""
    por: dict[tuple[str, int], list[Ocupacao]] = {}
    for o in cadastro:
        por.setdefault((o.tipo, o.numero), []).append(o)

    def resolve(rot: str, dia: date) -> str:
        m = _ROTULO_RE.match(rot)
        linhas = por.get((m[1].lower(), int(m[2]))) if m else None
        if not linhas:
            return rot
        if len({o.atendente for o in linhas}) == 1:
            o = linhas[0]
        else:
            o = next((x for x in linhas if (x.vig_inicio is None or dia >= x.vig_inicio)
                      and (x.vig_fim is None or dia <= x.vig_fim)), None)
            if o is None:
                return rot
        return f"{o.atendente} ({o.grupo}) — {rot}" if o.grupo else f"{o.atendente} — {rot}"
    return resolve


def le_cadastro_guiches(raw: bytes) -> list[Ocupacao] | None:
    """Planilha com aba 'Guichês' (Guichê, Grupo/Empresa, Atendente, …, Vigência início, Vigência fim)
    e, opcionalmente, aba de salas (Sala, Supervisor(a), Empresa). None se não for esse arquivo."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — não é xlsx válido
        return None
    aba = next((ws for ws in wb.worksheets if norm(ws.title).startswith("guiche")), None)
    if aba is None:
        return None
    linhas = list(aba.iter_rows(values_only=True))
    if not linhas or norm(linhas[0][0]) != "guiche" or "atendente" not in [norm(c) for c in linhas[0]]:
        return None
    cab = [norm(c) for c in linhas[0]]
    col = lambda *nomes: next((i for i, c in enumerate(cab) if any(c.startswith(n) for n in nomes)), None)  # noqa: E731
    ia, ig, ii, if_ = col("atendente"), col("grupo"), col("vigencia inicio"), col("vigencia fim")

    def dt(v):
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        return datetime.strptime(str(v).strip(), "%d/%m/%Y").date() if v not in (None, "") else None

    out = []
    for r in linhas[1:]:
        if not r or r[0] in (None, "") or ia is None or not r[ia]:
            continue
        out.append(Ocupacao("guiche", int(r[0]), str(r[ia]).strip(), str(r[ig] or "").strip() if ig is not None else "",
                            dt(r[ii]) if ii is not None else None, dt(r[if_]) if if_ is not None else None))
    salas = next((ws for ws in wb.worksheets if "sala" in norm(ws.title)), None)
    if salas is not None:
        for r in list(salas.iter_rows(values_only=True))[1:]:
            if not r or not r[0] or len(r) < 2 or not r[1]:
                continue
            for num in {int(n) for n in re.findall(r"\d+", str(r[0]))}:
                out.append(Ocupacao("sala", num, str(r[1]).strip(), str(r[2] or "").strip() if len(r) > 2 else "", None, None))
    wb.close()
    return out


# ---------- PDFs do SGA ----------
TITULOS = {"atendimentos concluidos": "atendimentos", "servicos executados": "servicos",
           "tempos medios por atendente": "tempos"}
_PERIODO_RE = re.compile(r"Per[ií]odo de (\d{2}/\d{2}/\d{4}) a (\d{2}/\d{2}/\d{4})")
_TMA_RE = re.compile(r"^(?P<nome>\D.*?)\s+(?P<n>\d+)\s+(?P<tme>\d{2}:\d{2}:\d{2})\s+(?P<tmd>\d{2}:\d{2}:\d{2})"
                     r"\s+(?P<tma>\d{2}:\d{2}:\d{2})\s+(?P<tot>\d{2}:\d{2}:\d{2})\s*$")
_SVC_RE = re.compile(r"^(?P<servico>.*\S)\s+(?P<n>\d+)$")
_DATA_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")


class SgaInvalido(ValueError):
    pass


@dataclass
class LeituraSga:
    conjunto: str
    de: date
    ate: date
    registros: list[dict]
    avisos: list[str]


def _br(s: str) -> date:
    return datetime.strptime(s, "%d/%m/%Y").date()


def identifica(texto: str) -> tuple[str, date, date] | None:
    """Conjunto e período a partir do cabeçalho do relatório; None se não for PDF do SGA."""
    linhas = [norm(x) for x in texto.splitlines()[:8]]
    conj = next((c for t, c in TITULOS.items() if t in linhas), None)
    m = _PERIODO_RE.search(texto)
    if not conj or not m:
        return None
    return conj, _br(m[1]), _br(m[2])


def tempos_do_texto(texto: str, dia: date) -> tuple[list[dict], list[str]]:
    regs, avisos, sem_nome = [], [], 0
    for ln in texto.splitlines():
        m = _TMA_RE.match(ln.strip())
        if not m:
            continue
        rot = rotulo(m["nome"])
        if not rot:
            sem_nome += 1
            continue
        if ignorado(rot):
            continue
        regs.append({"data": dia, "atendente_origem": rot, "atendente": rot, "n": int(m["n"]),
                     **{k: duracao(m[k]) for k in ("tme", "tmd", "tma", "tot")}})
    if sem_nome:
        avisos.append(f"{sem_nome} linha(s) sem usuário descartada(s)")
    return regs, avisos


def servicos_do_texto(texto: str, dia: date) -> tuple[list[dict], list[str]]:
    regs, dentro, total = [], False, None
    for ln in (x.strip() for x in texto.splitlines()):
        if norm(ln) == "servico total":
            dentro = True
            continue
        if ln.startswith("Total:"):
            total = int(re.sub(r"\D", "", ln) or 0)
            break
        m = _SVC_RE.match(ln) if dentro else None
        if m:
            regs.append({"data": dia, "servico": m["servico"], "qtd": int(m["n"])})
    avisos = []
    if total is not None and total != sum(r["qtd"] for r in regs):
        avisos.append(f"Soma dos serviços ({sum(r['qtd'] for r in regs)}) difere do Total do relatório ({total})")
    return regs, avisos


def _junta_quebras(s: str) -> str:
    """Célula quebrada pelo PDF: 'Seguro Auto(Novo/\\nRenovação)' → sem espaço após / ( -; demais viram espaço."""
    s = re.sub(r"(?<=[/(\-])\s*\n\s*", "", s)
    return re.sub(r"\s+", " ", s).strip()


def atendimentos_das_tabelas(tabelas: list[list[list]]) -> tuple[list[dict], list[str]]:
    """Tickets: linha de dados (Senha, Data, Chamada, Início, Fim, Duração, Permanência, Serviço
    Triado, Atendente) seguida de linhas auxiliares (Observação, Serviços realizados…). Usa a
    tabela (não o texto) porque 'Serviço Triado' quebra em duas linhas no texto."""
    regs, ruins = [], 0
    for t in tabelas:
        for r in t:
            if not r or len(r) < 9 or not r[0] or not _DATA_RE.match(str(r[1] or "").strip()):
                continue
            try:
                regs.append({
                    "data": _br(r[1].strip()), "hora": int(str(r[2]).strip()[:2]),
                    "servico": _junta_quebras(r[7] or ""),
                    "atendente_origem": rotulo(r[8] or ""), "atendente": rotulo(r[8] or ""),
                    "dur": duracao(r[5].strip()) if r[5] else None, "perm": duracao(r[6].strip()) if r[6] else None,
                })
            except (ValueError, AttributeError):
                ruins += 1
    return regs, [f"{ruins} ticket(s) ilegível(is) descartado(s)"] if ruins else []


def le_pdf_sga(raw: bytes) -> LeituraSga | None:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        if not pdf.pages:
            return None
        texto = "\n".join(p.extract_text() or "" for p in pdf.pages)
        ident = identifica(texto)
        if not ident:
            return None
        conj, de, ate = ident
        if conj == "atendimentos":
            regs, avisos = atendimentos_das_tabelas([t for p in pdf.pages for t in p.extract_tables()])
            return LeituraSga(conj, de, ate, regs, avisos)
    if de != ate:
        # TMA e serviços vêm agregados no período: não dá para distribuir por dia
        raise SgaInvalido(f"Relatório de {de:%d/%m/%Y} a {ate:%d/%m/%Y}: exporte um arquivo por dia "
                          "(o painel precisa do valor diário).")
    regs, avisos = (tempos_do_texto if conj == "tempos" else servicos_do_texto)(texto, de)
    return LeituraSga(conj, de, ate, regs, avisos)


def e_consolidado_protocolos(cabecalho: list) -> bool:
    """'Ouvidorias por Situação': um protocolo por linha, com nome do associado. Não é fonte
    do painel (os diários já contam as mensagens) e traz dado pessoal — não importar."""
    c = {norm(x) for x in cabecalho}
    return "protocolo" in c and ("situacao" in c or "datacadastro" in c)
