"""Leitura de arquivos em grades (lista de linhas) por aba/tabela.

Cada leitor devolve [{"nome": str, "linhas": list[list]}]. O reconhecimento de
cabeçalho e de conjunto acontece depois, em recognize.py, igual para todos os formatos.
"""
import csv
import io
import json
from datetime import datetime
from pathlib import Path


class FormatoNaoSuportado(Exception):
    pass


def _limpa(linhas: list[list]) -> list[list]:
    """Apara textos e remove só as linhas vazias do fim: as do meio ficam, para que os
    números de linha relatados nos erros batam com os da planilha."""
    out = [[c.strip() if isinstance(c, str) else c for c in r] for r in linhas]
    while out and all(c in (None, "") for c in out[-1]):
        out.pop()
    return out


def le_csv(raw: bytes) -> list[dict]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            txt = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    amostra = txt[:20000]
    try:
        delim = csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
    except csv.Error:
        delim = ";" if amostra.count(";") > amostra.count(",") else ","
    return [{"nome": "CSV", "linhas": _limpa(list(csv.reader(io.StringIO(txt), delimiter=delim)))}]


def le_xlsx(raw: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    abas = [{"nome": ws.title, "linhas": _limpa([list(r) for r in ws.iter_rows(values_only=True)])} for ws in wb.worksheets]
    wb.close()
    return [a for a in abas if a["linhas"]]


def le_xls(raw: bytes) -> list[dict]:
    import xlrd
    wb = xlrd.open_workbook(file_contents=raw)
    abas = []
    for sh in wb.sheets():
        linhas = []
        for i in range(sh.nrows):
            row = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(datetime(*xlrd.xldate_as_tuple(c.value, wb.datemode)))
                elif c.ctype == xlrd.XL_CELL_NUMBER and float(c.value).is_integer():
                    row.append(int(c.value))
                else:
                    row.append(c.value)
            linhas.append(row)
        abas.append({"nome": sh.name, "linhas": _limpa(linhas)})
    return [a for a in abas if a["linhas"]]


def le_pdf(raw: bytes) -> list[dict]:
    """Extrai as tabelas de todas as páginas e concatena as que têm o mesmo nº de colunas
    (relatórios do SGA repetem o cabeçalho a cada página; recognize.py descarta as repetições)."""
    import pdfplumber
    grupos: dict[int, list[list]] = {}
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for pg in pdf.pages:
            for t in pg.extract_tables():
                t = _limpa([[(c or "").replace("\n", " ") for c in r] for r in t])
                if t:
                    grupos.setdefault(len(t[0]), []).extend(t)
            if not grupos:  # PDF sem bordas de tabela: tenta pelo alinhamento do texto
                t = pg.extract_table({"vertical_strategy": "text", "horizontal_strategy": "text"})
                t = _limpa([[(c or "") for c in r] for r in t or []])
                if t:
                    grupos.setdefault(len(t[0]), []).extend(t)
    if not grupos:
        raise FormatoNaoSuportado("Nenhuma tabela encontrada no PDF (é um PDF digitalizado/imagem?)")
    return [{"nome": f"Tabela {n} colunas", "linhas": ls} for n, ls in sorted(grupos.items(), key=lambda g: -len(g[1]))]


def extrai_raw_legado(raw: bytes) -> dict | None:
    """Painel HTML autocontido antigo: devolve o objeto `const RAW = {...}` embutido."""
    txt = raw.decode("utf-8", errors="ignore")
    i = txt.find("const RAW = ")
    if i < 0:
        return None
    obj, _ = json.JSONDecoder().raw_decode(txt, i + len("const RAW = "))
    return obj if isinstance(obj, dict) and "dias" in obj else None


LEITORES = {".csv": le_csv, ".txt": le_csv, ".xlsx": le_xlsx, ".xlsm": le_xlsx, ".xls": le_xls, ".pdf": le_pdf}


def le(nome: str, raw: bytes) -> list[dict]:
    ext = Path(nome).suffix.lower()
    if ext not in LEITORES:
        raise FormatoNaoSuportado(f"Formato {ext or '(sem extensão)'} não suportado. Use CSV, XLSX, XLS ou PDF.")
    return LEITORES[ext](raw)
